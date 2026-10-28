"""Printable facilitator guide for the Wealth Symposium activity."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

RED = RGBColor(0xDB, 0x00, 0x11)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x4A, 0x4A, 0x4A)
LINE = RGBColor(0xE4, 0xE4, 0xE4)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
WASH = RGBColor(0xF8, 0xF4, 0xF4)
FONT = "Calibri"

W = Inches(8.27)
H = Inches(11.69)


def set_run(run, size, bold=False, color=INK, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    # PowerPoint on Windows uses Calibri. Pin the latin typeface so it does not fall back.
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = rpr.makeelement(qn(tag), {})
            rpr.append(node)
        node.set("typeface", font)


def add_text(slide, l, t, w, h, paragraphs, size=13, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(l, t, w, h)
    shape.word_wrap = True
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.anchor = anchor
    for index, item in enumerate(paragraphs):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        if isinstance(item, str):
            item = {"text": item}
        p.alignment = item.get("align", align)
        p.space_before = Pt(item.get("before", 0))
        p.space_after = Pt(item.get("after", 2))
        p.line_spacing = item.get("leading", 1.05)
        run = p.add_run()
        run.text = item["text"]
        set_run(run, item.get("size", size), item.get("bold", bold), item.get("color", color))
    return shape


def rect(slide, l, t, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def page(prs, section):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, PAPER)
    rect(slide, 0, 0, Inches(0.12), H, RED)
    add_text(
        slide,
        Inches(0.42),
        Inches(0.28),
        Inches(5.6),
        Inches(0.28),
        [{"text": "WEALTH SYMPOSIUM  ·  FACILITATOR GUIDE", "size": 11, "bold": True, "color": RED}],
    )
    add_text(
        slide,
        Inches(5.5),
        Inches(0.28),
        Inches(2.35),
        Inches(0.28),
        [{"text": section, "size": 11, "color": MUTED, "align": PP_ALIGN.RIGHT}],
    )
    rect(slide, Inches(0.42), Inches(0.62), Inches(7.42), Inches(0.015), LINE)
    add_text(
        slide,
        Inches(0.42),
        Inches(11.28),
        Inches(7.42),
        Inches(0.24),
        [{"text": "AI Powered, Human Led  ·  Do not hand this sheet to the groups", "size": 10, "color": MUTED}],
    )
    return slide


def h1(slide, text, top):
    add_text(slide, Inches(0.42), top, Inches(7.42), Inches(0.42), [{"text": text, "size": 26, "bold": True}])


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "Wealth Symposium facilitator guide"
    prs.core_properties.subject = "AI Powered, Human Led"

    s = page(prs, "01  Purpose")
    h1(s, "Run the hour. Don't design it for them.", Inches(0.78))
    add_text(
        s,
        Inches(0.42),
        Inches(1.32),
        Inches(7.42),
        Inches(1.15),
        [
            {
                "text": "Ten groups. Five relationship managers. One laptop, driven by one leader. In thirty minutes of making they turn one real painpoint into a named agent, three screens, and a clickable demo. Then each group has ninety seconds.",
                "size": 15,
                "after": 8,
            },
            {
                "text": "You are not teaching a tool, and you are not picking a winner. You are there to see which pain they choose when the choice is free, and whether the agent they describe actually stays inside that pain.",
                "size": 15,
            },
        ],
    )

    blocks = [
        ("Your job", "Keep every group inside one painpoint, and get every group to the ninety seconds with something they can show. A rough demo that they can walk through beats a polished prompt that never gets built."),
        ("A good group looks like", "They argue about the pain, not about the wording. The agent does one job. The three screens are three different moments, in order. The benefit is something a client or an RM would notice on Thursday, not a slogan."),
        ("What you take away", "A tally of which painpoints were chosen, including how many groups wrote their own. The agent name and the one benefit, in their words. Do not rank the screens."),
    ]
    top = 2.7
    for title, body in blocks:
        rect(s, Inches(0.42), Inches(top), Inches(0.08), Inches(1.15), RED)
        add_text(
            s,
            Inches(0.66),
            Inches(top),
            Inches(7.15),
            Inches(1.15),
            [
                {"text": title, "size": 16, "bold": True, "after": 4},
                {"text": body, "size": 14, "color": MUTED},
            ],
        )
        top += 1.28

    add_text(
        s,
        Inches(0.42),
        Inches(6.7),
        Inches(7.42),
        Inches(1.5),
        [
            {"text": "Three things, if the room gets noisy", "size": 16, "bold": True, "after": 6},
            {"text": "1   One leader. Everyone else talks. They press Start fresh, not Resume.", "size": 14, "after": 3},
            {"text": "2   Stay on the pain they picked. If the sentence could belong to any agent, it is not done.", "size": 14, "after": 3},
            {"text": "3   Stop polishing. One or two edits on the demo, then present.", "size": 14},
        ],
    )

    rect(s, Inches(0.42), Inches(8.45), Inches(7.42), Inches(2.5), WASH)
    add_text(
        s,
        Inches(0.62),
        Inches(8.6),
        Inches(7.05),
        Inches(2.2),
        [
            {"text": "Leave the choice free", "size": 16, "bold": True, "after": 6},
            {
                "text": "Do not steer a group off a card because another group already took it. If seven groups pick the same pain, that is the finding. Say so in the close. Do not apologise for it, and do not ask them to pick again so the wall looks balanced.",
                "size": 14,
                "after": 8,
            },
            {
                "text": "The only time you intervene on the choice is when they have drifted into a feature request. “An agent that summarises everything” is not a painpoint. Send them back to the card, or to a situation they can picture with a client in the room.",
                "size": 14,
            },
        ],
    )

    s = page(prs, "02  Room and clock")
    h1(s, "Set the room before they sit down.", Inches(0.78))
    add_text(
        s,
        Inches(0.42),
        Inches(1.28),
        Inches(7.42),
        Inches(2.35),
        [
            {"text": "On each laptop, before the session", "size": 16, "bold": True, "after": 4},
            {"text": "Open the activity in one browser tab, full screen. Use the link you have already tested on that network.", "size": 14, "after": 2},
            {"text": "https://sync710n.github.io/wealthsymposium26/", "size": 14, "bold": True, "color": RED, "after": 4},
            {"text": "Leave it on the cover. Do not press Start fresh until you say go. If a laptop was used in rehearsal, Start fresh once, then return to the cover only if Resume is not showing. You want a clean cover: Start fresh, and no Resume.", "size": 14, "after": 4},
            {"text": "Use the zoom control until a person at the back of the table can read a painpoint card. The printed desk card stays face up. You keep this guide. They do not get it.", "size": 14, "after": 4},
            {"text": "Your laptop is the only one on the big screen. Have one finished example already generated, in case a live generation stalls.", "size": 14},
        ],
    )

    add_text(s, Inches(0.42), Inches(3.75), Inches(7.42), Inches(0.32), [{"text": "The clock on their laptop", "size": 16, "bold": True}])
    add_text(
        s,
        Inches(0.42),
        Inches(4.08),
        Inches(7.42),
        Inches(0.7),
        [
            {
                "text": "Each phase has its own timer. It starts when that laptop enters the phase, not when you start speaking. Call the changes out loud so the room stays together. A group that finishes early may go on. At your cutoff, everyone moves, even if Next is the next thing they do.",
                "size": 13,
                "color": MUTED,
            }
        ],
    )

    rows = [
        ("Phase", "On screen", "Their timer", "Your cutoff"),
        ("0", "Painpoint", "5:00", "Minute 0 to 5"),
        ("1", "Idea brief", "5:00", "Minute 5 to 10"),
        ("2", "Design 3 screens", "8:00", "Minute 10 to 18"),
        ("3", "Demo", "12:00", "Minute 18 to 30"),
    ]
    table_top = 4.85
    row_h = 0.38
    widths = [0.7, 2.3, 1.5, 2.9]
    left = 0.42
    for r, row in enumerate(rows):
        x = left
        for c, cell in enumerate(row):
            if r == 0:
                rect(s, Inches(x), Inches(table_top + r * row_h), Inches(widths[c]), Inches(row_h), RED)
                color, bold = PAPER, True
            else:
                if r % 2 == 0:
                    rect(s, Inches(x), Inches(table_top + r * row_h), Inches(widths[c]), Inches(row_h), WASH)
                color, bold = INK, c == 0
            add_text(
                s,
                Inches(x + 0.08),
                Inches(table_top + r * row_h + 0.06),
                Inches(widths[c] - 0.1),
                Inches(0.28),
                [{"text": cell, "size": 13, "bold": bold, "color": color}],
            )
            x += widths[c]

    add_text(s, Inches(0.42), Inches(6.95), Inches(7.42), Inches(0.32), [{"text": "The hour around the thirty minutes", "size": 16, "bold": True}])
    run = [
        ("3 min", "Cover stays up. Nominate the leader. Say what the hour is for."),
        ("6 min", "You demo the clicks. Do not invent their pain for them."),
        ("30 min", "They make. You walk. You call the four cutoffs."),
        ("2 min", "Everyone opens Present on Screen 1."),
        ("15 min", "Ten groups, ninety seconds. You keep time. No critique in between."),
        ("3 min", "Read the tally. Name the pains that came up. Stop."),
    ]
    y = 7.35
    for time, text in run:
        add_text(s, Inches(0.42), Inches(y), Inches(1.15), Inches(0.36), [{"text": time, "size": 13, "bold": True, "color": RED}])
        add_text(s, Inches(1.6), Inches(y), Inches(6.2), Inches(0.4), [{"text": text, "size": 13}])
        y += 0.38

    add_text(
        s,
        Inches(0.42),
        Inches(9.75),
        Inches(7.42),
        Inches(1.3),
        [
            {"text": "If the slot is 45 minutes", "size": 16, "bold": True, "after": 4},
            {
                "text": "Cut the demo to four minutes and use the example you already built. Keep the thirty minutes of making. Hear two groups in full, then ask every other group for only the painpoint and the agent name. Collect the rest of the tally from the screens as you walk. Do not shrink the making time to fit more presentations.",
                "size": 13,
                "color": MUTED,
            },
        ],
    )

    s = page(prs, "03  Open and demo")
    h1(s, "Say less than you want to.", Inches(0.78))
    add_text(
        s,
        Inches(0.42),
        Inches(1.28),
        Inches(7.42),
        Inches(2.15),
        [
            {"text": "Open, while they are still on the cover", "size": 16, "bold": True, "after": 6},
            {
                "text": "“You are not here to learn a tool. You are here to design the assistant you wish you had beside you with a client. One person drives the laptop. The other four decide. Pick the pain that is actually painful, not the one that sounds impressive. At the end you have ninety seconds: the pain, the agent, the demo, and one benefit.”",
                "size": 14,
                "after": 8,
            },
            {
                "text": "Then point at the cover. “When I say go, the leader presses Start fresh. Not Resume. Resume is only if this laptop already started and died.”",
                "size": 14,
            },
        ],
    )

    add_text(s, Inches(0.42), Inches(3.55), Inches(7.42), Inches(0.32), [{"text": "The live demo, on your laptop only", "size": 16, "bold": True}])
    add_text(
        s,
        Inches(0.42),
        Inches(3.95),
        Inches(7.42),
        Inches(0.55),
        [
            {
                "text": "Walk the clicks. Do not workshop the content. Use a pain you are willing to throw away. Tell them that. If a generation spinner sits there, talk over it once, then switch to the example you rehearsed.",
                "size": 13,
                "color": MUTED,
            }
        ],
    )
    demo = [
        ("Cover", "Show Start fresh and Resume. Say which one they will press."),
        ("Phase 0", "Click one card. Read the title, not the whole paragraph. Press Next."),
        ("Phase 1", "The pain stays locked at the top. Click a box and type. Show that a suggestion does nothing until Accept."),
        ("Phase 2", "Three prompts, in order. Generate screen sits on the picture. Edit the words after, and that picture is out of date until you generate it again."),
        ("Phase 3", "Generate the demo. It writes one screen, then the next. Show Present, and the four things they will say."),
    ]
    y = 4.6
    for label, text in demo:
        rect(s, Inches(0.42), Inches(y), Inches(1.35), Inches(0.85), WASH)
        add_text(
            s,
            Inches(0.5),
            Inches(y + 0.22),
            Inches(1.2),
            Inches(0.45),
            [{"text": label, "size": 13, "bold": True, "color": RED}],
        )
        add_text(s, Inches(1.9), Inches(y + 0.08), Inches(5.9), Inches(0.75), [{"text": text, "size": 13}])
        y += 0.95

    add_text(
        s,
        Inches(0.42),
        Inches(9.5),
        Inches(7.42),
        Inches(1.5),
        [
            {"text": "Then get out of the content", "size": 16, "bold": True, "after": 4},
            {
                "text": "“That pain was mine, for the demo. Yours starts now. Press Start fresh.” Do not leave your example on the big screen while they choose. Turn the projector to the cover, or off. Otherwise the room designs your agent instead of theirs.",
                "size": 14,
            },
        ],
    )

    s = page(prs, "04  While they work")
    h1(s, "Walk. Don't take the laptop.", Inches(0.78))
    add_text(
        s,
        Inches(0.42),
        Inches(1.22),
        Inches(7.42),
        Inches(0.55),
        [
            {
                "text": "Visit every table in each phase. Ask one question, then leave. If you stay, you become the fifth designer and the group waits for you.",
                "size": 14,
                "color": MUTED,
            }
        ],
    )

    phases = [
        (
            "0   Painpoint",
            "5 minutes",
            "One card, or their own sentence in the bottom row. Next stays grey until then.",
            "Choosing a pain, or already designing the agent. A pain is a situation. An idea is a feature.",
            "“Which of these would you want gone before Thursday?”",
        ),
        (
            "1   Idea brief",
            "5 minutes",
            "A short name and three boxes, in their words. Next needs all four. The pain at the top is locked.",
            "A name that sounds like a product. A job that is “helps with everything.” A benefit that fits any agent.",
            "“Say the name out loud.” For a thin box: “Write a bit more. After about ten words, Refine can help. You still press Accept.”",
        ),
        (
            "2   Three screens",
            "8 minutes",
            "Three different moments, in order, each with a picture. Next stays grey if one is missing or out of date.",
            "Three versions of the same moment. A prompt that restates the pain. A group waiting on one perfect picture.",
            "“Screen 1 is the assistant starting. Screen 3 is you doing something you could not do then. Regenerate that screen only.”",
        ),
        (
            "3   Demo",
            "12 minutes",
            "Generate the demo, then one or two edits. Present is how they will show it.",
            "Still rewriting the brief. Refreshing the page. Asking for a fifth feature. Silent, watching a spinner.",
            "“One change, then stop. If this died, press Resume. Start fresh deletes the work.”",
        ),
    ]
    y = 1.85
    for title, time, doing, watch, line in phases:
        rect(s, Inches(0.42), Inches(y), Inches(7.42), Inches(2.15), WASH)
        add_text(s, Inches(0.55), Inches(y + 0.06), Inches(5.4), Inches(0.3), [{"text": title, "size": 15, "bold": True}])
        add_text(
            s,
            Inches(5.7),
            Inches(y + 0.08),
            Inches(2.0),
            Inches(0.28),
            [{"text": time, "size": 13, "bold": True, "color": RED, "align": PP_ALIGN.RIGHT}],
        )
        add_text(
            s,
            Inches(0.55),
            Inches(y + 0.38),
            Inches(7.15),
            Inches(1.7),
            [
                {"text": "They need  " + doing, "size": 12, "after": 2},
                {"text": "You watch  " + watch, "size": 12, "after": 2},
                {"text": "If stuck  " + line, "size": 12},
            ],
        )
        y += 2.25

    s = page(prs, "05  Ninety seconds")
    h1(s, "Chair it. Don't coach it.", Inches(0.78))
    add_text(
        s,
        Inches(0.42),
        Inches(1.25),
        Inches(7.42),
        Inches(0.85),
        [
            {
                "text": "Same order every time, so the room can compare. You keep a hand up at seventy-five seconds and you thank them when the ninety is done. No questions from the floor. No “what I would have done” from you. That conversation happens after, if at all.",
                "size": 14,
            }
        ],
    )
    beats = [
        ("1", "The pain", "The situation, not the solution. One or two sentences."),
        ("2", "The agent", "The name, and the one job it does."),
        ("3", "The demo", "Screen 1, then 2, then 3. What the assistant does, and what the RM still decides."),
        ("4", "The benefit", "One change for the RM or the client. Then stop."),
    ]
    y = 2.25
    for n, title, body in beats:
        rect(s, Inches(0.42), Inches(y), Inches(0.48), Inches(0.48), RED)
        add_text(
            s,
            Inches(0.42),
            Inches(y + 0.06),
            Inches(0.48),
            Inches(0.38),
            [{"text": n, "size": 16, "bold": True, "color": PAPER, "align": PP_ALIGN.CENTER}],
        )
        add_text(s, Inches(1.05), Inches(y), Inches(6.7), Inches(0.28), [{"text": title, "size": 15, "bold": True}])
        add_text(s, Inches(1.05), Inches(y + 0.3), Inches(6.7), Inches(0.28), [{"text": body, "size": 13, "color": MUTED}])
        y += 0.72

    add_text(s, Inches(0.42), Inches(5.2), Inches(7.42), Inches(0.3), [{"text": "Tally, while they talk", "size": 16, "bold": True}])
    add_text(
        s,
        Inches(0.42),
        Inches(5.52),
        Inches(7.42),
        Inches(0.4),
        [{"text": "One mark per group. Write the agent name in the margin if you can hear it. This is the page you close on.", "size": 12, "color": MUTED}],
    )
    pains = [
        "01  Unclear client priorities",
        "02  Scattered information, lengthy preparation",
        "03  Too much to juggle in meetings",
        "04  Competing client goals",
        "05  Lengthy comparisons and explanations",
        "06  Lengthy onboarding, KYC and EDD",
        "07  Lengthy discount and bespoke approvals",
        "08  Complex cross-border requirements",
        "09  Too much market information to absorb",
        "Own  They wrote their own",
    ]
    y = 6.0
    for title in pains:
        rect(s, Inches(0.42), Inches(y), Inches(7.42), Inches(0.38), WASH if pains.index(title) % 2 == 0 else PAPER)
        add_text(s, Inches(0.52), Inches(y + 0.05), Inches(5.3), Inches(0.28), [{"text": title, "size": 13}])
        add_text(s, Inches(5.9), Inches(y + 0.05), Inches(1.8), Inches(0.28), [{"text": "□ □ □ □ □ □ □ □ □ □", "size": 12, "color": MUTED}])
        y += 0.38

    add_text(
        s,
        Inches(0.42),
        Inches(9.95),
        Inches(7.42),
        Inches(1.15),
        [
            {"text": "The close, after the last group", "size": 16, "bold": True, "after": 4},
            {
                "text": "Read what the marks show. “Preparation came up four times. The meeting itself came up three. Two groups wrote their own.” Then stop. Do not award a winner, and do not promise which agent will get built.",
                "size": 14,
            },
        ],
    )

    s = page(prs, "06  If it breaks")
    h1(s, "Recover quietly.", Inches(0.78))
    add_text(
        s,
        Inches(0.42),
        Inches(1.25),
        Inches(7.42),
        Inches(0.55),
        [
            {
                "text": "Groups should not see the debug menu. Double-click the words Wealth Symposium only on your own laptop, and only if you need it.",
                "size": 14,
                "color": MUTED,
            }
        ],
    )
    fixes = [
        ("Resume is showing and they have not started", "Someone opened this laptop earlier. Press Start fresh before the group sits down. Say nothing about it."),
        ("The page froze or they refreshed", "Press Resume. Not Start fresh. Start fresh clears the pain, the brief, the screens, and the demo."),
        ("Next is grey on the screens", "A picture is missing, or a prompt was edited after the picture was made. Hover that picture and press Generate screen again. Only that one."),
        ("The picture failed", "Stay on the page and press Generate screen again. Don't rewrite the whole story because one image failed."),
        ("The demo is taking a while", "That is normal. It writes one screen, then the next. Ask them what they will say for Screen 1 while they wait. Do not refresh."),
        ("They cannot read the cards", "Use the zoom control. Don't ask them to lean in for thirty minutes."),
        ("The assistant wrote something they hate", "Dismiss, don't Accept. It does not change the box until they accept it."),
        ("One person is typing everyone's idea", "“Leader drives. The other four decide the next sentence.” Then move to the next table."),
        ("You are tempted to fix the words", "Don't. A clumsy sentence they believe is more useful to you than a clean sentence you wrote."),
    ]
    y = 1.95
    for title, body in fixes:
        add_text(s, Inches(0.42), Inches(y), Inches(7.42), Inches(0.26), [{"text": title, "size": 14, "bold": True}])
        add_text(s, Inches(0.42), Inches(y + 0.26), Inches(7.42), Inches(0.48), [{"text": body, "size": 13, "color": MUTED}])
        y += 0.82

    add_text(
        s,
        Inches(0.42),
        Inches(9.5),
        Inches(7.42),
        Inches(1.5),
        [
            {"text": "After the room clears", "size": 16, "bold": True, "after": 4},
            {
                "text": "Photograph the tally before you pack. If a laptop still has a demo up, that is the group's record of the session. Don't press Start fresh on the way out.",
                "size": 14,
            },
        ],
    )

    out = "/workspace/docs/Wealth-Symposium-Facilitator-Guide.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
